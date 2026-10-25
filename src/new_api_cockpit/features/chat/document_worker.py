"""Isolated binary parsing entry point. stdin is bytes; stdout is bounded JSON.

Never evaluate Office formulas/macros, PDF JavaScript, embedded objects or links.
No application configuration is imported and no inherited credentials are used.
"""

import base64
from contextlib import closing
import io
import json
import os
from pathlib import Path
import posixpath
import resource
import subprocess
import sys
import warnings
import zipfile

# -I deliberately ignores cwd/PYTHONPATH; only this installed package is trusted.
sys.path.insert(0, str(Path(__file__).resolve().parents[3]))
from new_api_cockpit.features.chat import config  # noqa: E402


class LimitError(Exception):
    pass


def bounded_text(lines):
    size, result = 0, []
    for line in lines:
        line = str(line).replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
        # Extracted controls are inert but must be valid portable prompt text.
        line = "".join(c for c in line if ord(c) >= 32 or c in "\n\t")
        size += len(line.encode()) + 1
        if size > config.DOCUMENT_TEXT_BYTES:
            raise LimitError()
        result.append(line)
    return "\n".join(result)


def image(raw, expected=None):
    from PIL import Image, ImageOps

    Image.MAX_IMAGE_PIXELS = config.IMAGE_MAX_PIXELS
    warnings.simplefilter("error", Image.DecompressionBombWarning)
    with Image.open(io.BytesIO(raw)) as source:
        if expected and source.format not in expected:
            raise ValueError()
        if (
            source.width * source.height > config.IMAGE_MAX_PIXELS
            or getattr(source, "n_frames", 1) != 1
        ):
            raise LimitError()
        source.load()
        frame = ImageOps.exif_transpose(source)
        frame.thumbnail((config.IMAGE_MAX_EDGE, config.IMAGE_MAX_EDGE))
        if frame.mode in {"RGBA", "LA"} or "transparency" in frame.info:
            rgba = frame.convert("RGBA")
            rgb = Image.new("RGB", rgba.size, "white")
            rgb.paste(rgba, mask=rgba.getchannel("A"))
        else:
            rgb = frame.convert("RGB")
        out = io.BytesIO()
        rgb.save(out, format="JPEG", quality=85)
        data = out.getvalue()
        if len(data) > config.IMAGE_ENCODED_BYTES:
            raise LimitError()
        return {"mime": "image/jpeg", "data": base64.b64encode(data).decode("ascii")}


def office_zip(raw):
    archive = zipfile.ZipFile(io.BytesIO(raw))
    entries = archive.infolist()
    names = [i.filename for i in entries]
    if (
        len(entries) > 4096
        or len(set(names)) != len(names)
        or sum(i.file_size for i in entries) > config.DOCUMENT_ZIP_BYTES
    ):
        raise LimitError()
    for entry in entries:
        if (
            entry.flag_bits & 1
            or entry.filename.startswith(("/", "\\"))
            or ".." in entry.filename.split("/")
            or "\\" in entry.filename
        ):
            raise ValueError()
        if entry.file_size > config.DOCUMENT_XML_BYTES and entry.filename.endswith(
            (".xml", ".rels")
        ):
            raise LimitError()
        if entry.file_size > max(1048576, entry.compress_size * 200):
            raise LimitError()
    if "[Content_Types].xml" not in names:
        raise ValueError()
    return archive


def pdf(raw):
    import pypdfium2 as pdfium

    if not raw.startswith(b"%PDF-"):
        raise ValueError()
    lines, pictures = [], []
    with pdfium.PdfDocument(raw) as doc:
        if not 1 <= len(doc) <= config.DOCUMENT_MAX_PAGES:
            raise LimitError()
        for number in range(len(doc)):
            with closing(doc[number]) as page:
                with closing(page.get_textpage()) as textpage:
                    if textpage.count_chars() > config.DOCUMENT_TEXT_BYTES:
                        raise LimitError()
                    text = textpage.get_text_range()
                lines.extend([f"[PDF 第 {number + 1} 页]", text])
                if not text.strip():
                    if len(pictures) >= config.IMAGE_COUNT_MAX:
                        raise LimitError()
                    width, height = page.get_size()
                    if min(width, height) <= 0:
                        raise ValueError()
                    scale = min(2.0, config.IMAGE_MAX_EDGE / max(width, height))
                    bitmap = page.render(scale=scale)
                    try:
                        pil = bitmap.to_pil()
                        out = io.BytesIO()
                        pil.save(out, format="PNG")
                        pictures.append(image(out.getvalue(), {"PNG"}))
                    finally:
                        bitmap.close()
                    lines.append("本页为扫描/图像页，已附页图；请结合图像读取。")
                # Bound accumulated extraction before moving to another page.
                if sum(len(t.encode()) for t in lines) > config.DOCUMENT_TEXT_BYTES:
                    raise LimitError()
    return {"text": bounded_text(lines), "images": pictures}


def word(raw):
    from defusedxml import ElementTree as ET

    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    r = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    a = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    lines, pictures = [], []
    with office_zip(raw) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
        if root.tag != w + "document":
            raise ValueError()

        def paragraph(node):
            return "".join(
                n.text or ""
                if n.tag == w + "t"
                else "\t"
                if n.tag == w + "tab"
                else "\n"
                if n.tag in {w + "br", w + "cr"}
                else ""
                for n in node.iter()
            )

        for node in root.find(w + "body"):
            if node.tag == w + "p":
                lines.append(paragraph(node))
            elif node.tag == w + "tbl":
                for row in node.findall(w + "tr"):
                    lines.append(
                        "\t".join(
                            " / ".join(paragraph(p) for p in cell.iter(w + "p"))
                            for cell in row.findall(w + "tc")
                        )
                    )
        for name in sorted(archive.namelist()):
            if name.startswith(("word/header", "word/footer")) and name.endswith(
                ".xml"
            ):
                lines.extend(
                    paragraph(p)
                    for p in ET.fromstring(archive.read(name)).iter(w + "p")
                )
        text = bounded_text(lines)
        if "word/_rels/document.xml.rels" in archive.namelist():
            rels = {
                n.attrib["Id"]: n.attrib
                for n in ET.fromstring(archive.read("word/_rels/document.xml.rels"))
            }
            seen = set()
            for blip in root.iter(a + "blip"):
                rel = rels.get(blip.get(r + "embed"), {})
                if not rel or rel.get("TargetMode") == "External":
                    continue
                name = posixpath.normpath(posixpath.join("word", rel["Target"]))
                if not name.startswith("word/media/") or name in seen:
                    continue
                seen.add(name)
                if Path(name).suffix.lower() not in {
                    ".png",
                    ".jpg",
                    ".jpeg",
                    ".webp",
                    ".gif",
                    ".bmp",
                }:
                    text = bounded_text(
                        [text, "文档包含未解析的嵌入图形；未将其当作已读取图像。"]
                    )
                    continue
                if len(pictures) >= config.IMAGE_COUNT_MAX:
                    raise LimitError()
                pictures.append(image(archive.read(name)))
    return {"text": text, "images": pictures}


def excel(raw, extension):
    lines, cells = [], 0
    if extension == ".xlsx":
        from openpyxl import load_workbook

        with office_zip(raw) as archive:
            if "xl/workbook.xml" not in archive.namelist():
                raise ValueError()
        book = load_workbook(
            io.BytesIO(raw), read_only=True, data_only=False, keep_links=False
        )
        cache = load_workbook(
            io.BytesIO(raw), read_only=True, data_only=True, keep_links=False
        )
        try:
            if len(book.worksheets) > 64:
                raise LimitError()
            for sheet, values in zip(book.worksheets, cache.worksheets):
                if (sheet.max_row or 0) * (
                    sheet.max_column or 0
                ) > config.DOCUMENT_MAX_CELLS:
                    raise LimitError()
                # Do not silently omit cells when a third-party writer supplied
                # an incorrect worksheet dimension. Count actual streamed rows.
                sheet.reset_dimensions()
                values.reset_dimensions()
                lines.append(f"[Excel 工作表：{sheet.title}]")
                for i, (row, cached) in enumerate(
                    zip(
                        sheet.iter_rows(values_only=True),
                        values.iter_rows(values_only=True),
                    ),
                    1,
                ):
                    if i > config.DOCUMENT_MAX_CELLS:
                        raise LimitError()
                    cells += len(row)
                    if cells > config.DOCUMENT_MAX_CELLS:
                        raise LimitError()
                    result = []
                    for value, cv in zip(row, cached):
                        if isinstance(value, str) and value.startswith("="):
                            value += (
                                f" [缓存值：{cv if cv is not None else '尚未计算'}]"
                            )
                        result.append("" if value is None else str(value))
                    if any(result):
                        lines.append(str(i) + "\t" + "\t".join(result))
                    if sum(len(t) for t in lines[-2:]) > config.DOCUMENT_TEXT_BYTES:
                        raise LimitError()
        finally:
            book.close()
            cache.close()
    else:
        import xlrd

        if not raw.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
            raise ValueError()
        with xlrd.open_workbook(file_contents=raw, on_demand=True) as book:
            if book.nsheets > 64:
                raise LimitError()
            for sheet in book.sheets():
                cells += sheet.nrows * sheet.ncols
                if cells > config.DOCUMENT_MAX_CELLS:
                    raise LimitError()
                lines.append(f"[Excel 工作表：{sheet.name}]")
                for i in range(sheet.nrows):
                    values = []
                    for cell in sheet.row(i):
                        value = (
                            xlrd.xldate.xldate_as_datetime(
                                cell.value, book.datemode
                            ).isoformat()
                            if cell.ctype == xlrd.XL_CELL_DATE
                            else cell.value
                        )
                        values.append(str(value))
                    lines.append(str(i + 1) + "\t" + "\t".join(values))
    return {"text": bounded_text(lines), "images": []}


def parse(raw, extension):
    if extension == ".pdf":
        return pdf(raw)
    if extension == ".docx":
        return word(raw)
    if extension in {".xlsx", ".xls"}:
        return excel(raw, extension)
    if extension == ".doc":
        if not raw.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
            raise ValueError()
        if not os.path.isfile("/usr/bin/antiword"):
            raise ImportError()
        result = subprocess.run(
            ["/usr/bin/antiword", "-m", "UTF-8.txt", "-w", "0", "-"],
            input=raw,
            capture_output=True,
            timeout=10,
            check=True,
        )
        return {"text": bounded_text([result.stdout.decode("utf-8")]), "images": []}
    expected = {
        ".png": {"PNG"},
        ".jpg": {"JPEG"},
        ".jpeg": {"JPEG"},
        ".webp": {"WEBP"},
        ".gif": {"GIF"},
        ".bmp": {"BMP"},
    }[extension]
    return {
        "text": "图片已读取，请结合随附图像分析。",
        "images": [image(raw, expected)],
    }


def main():
    # Applied inside the fresh single-threaded process, never preexec_fn in the
    # request pool's multithreaded parent. Limits also cover antiword children.
    if sys.platform == "linux":
        resource.setrlimit(resource.RLIMIT_AS, (256 * 1024 * 1024,) * 2)
    resource.setrlimit(resource.RLIMIT_CPU, (15, 16))
    resource.setrlimit(resource.RLIMIT_FSIZE, (config.MODEL_REQUEST_BYTES * 2,) * 2)
    resource.setrlimit(resource.RLIMIT_CORE, (0, 0))
    try:
        raw = sys.stdin.buffer.read(config.UPLOAD_FILE_BYTES + 1)
        if len(raw) > config.UPLOAD_FILE_BYTES:
            raise LimitError()
        result = parse(raw, sys.argv[1])
        data = json.dumps(result, ensure_ascii=False).encode()
        if len(data) > config.MODEL_REQUEST_BYTES * 2:
            raise LimitError()
    except ImportError:
        data = b'{"error":"CHAT_FILE_PARSER_UNAVAILABLE"}'
    except (MemoryError, LimitError):
        data = b'{"error":"CHAT_FILE_PARSE_LIMIT"}'
    except Exception:
        data = b'{"error":"CHAT_FILE_INVALID"}'
    sys.stdout.buffer.write(data)


if __name__ == "__main__":
    main()
