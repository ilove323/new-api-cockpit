"""Credential-free, bounded render process. No model code, URLs or user paths."""

import io
import json
import math
from pathlib import Path
import re
import resource
import sys
from xml.sax.saxutils import escape

MAX_BYTES = 4 * 1024 * 1024


def font_path():
    # Fixed administrator-installed fonts only; never a model-supplied path.
    for name in (
        "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
    ):
        if Path(name).is_file():
            return name
    raise ValueError("Chinese font unavailable")


def fields(value, required, optional=()):
    if (
        not isinstance(value, dict)
        or not set(required) <= value.keys()
        or value.keys() - set(required) - set(optional)
    ):
        raise ValueError("Invalid fields")


def string(value, maximum=2000):
    if (
        not isinstance(value, str)
        or len(value) > maximum
        or any(ord(c) < 32 and c not in "\n\t" for c in value)
    ):
        raise ValueError("Invalid text")
    return value


def number(value, low, high, *, integer=False):
    if (
        type(value) not in (int, float)
        or not math.isfinite(value)
        or not low <= value <= high
        or (integer and type(value) is not int)
    ):
        raise ValueError("Invalid number")
    return value


def blocks(content):
    lines = content.splitlines()
    result, index, code = [], 0, False
    while index < len(lines):
        line = lines[index]
        index += 1
        if line.startswith("```"):
            code = not code
            continue
        if not line.strip():
            continue
        if (
            not code
            and "|" in line
            and index < len(lines)
            and re.fullmatch(
                r"\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*", lines[index]
            )
        ):
            rows = [[c.strip() for c in line.strip().strip("|").split("|")]]
            index += 1
            while index < len(lines) and "|" in lines[index] and lines[index].strip():
                rows.append(
                    [c.strip() for c in lines[index].strip().strip("|").split("|")]
                )
                index += 1
            if (
                len(rows) > 1000
                or not 1 <= len(rows[0]) <= 8
                or any(
                    len(row) != len(rows[0]) or any(len(c) > 200 for c in row)
                    for row in rows
                )
            ):
                raise ValueError("Table limits")
            result.append(("table", rows))
            continue
        match = None if code else re.match(r"^(#{1,6})\s+(.+)$", line)
        if match:
            result.append(("h" + str(len(match[1])), match[2]))
        else:
            line = re.sub(r"^\s*[-*+]\s+", "• ", line) if not code else line
            # Bounded paragraphs prevent a single unsplittable flowable.
            result.extend(
                ("code" if code else "p", line[i : i + 800])
                for i in range(0, len(line), 800)
            )
    if len(result) > 3000:
        raise ValueError("Too many blocks")
    return result or [("p", "")]


def inline(text):
    # Escape first: ReportLab markup never receives model XML, links or images.
    return re.sub(r"\*\*([^*]+)\*\*", r"<b>\1</b>", escape(text)).replace("\t", "    ")


def pdf(content):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (
        SimpleDocTemplate,
        Paragraph,
        Spacer,
        LongTable,
        TableStyle,
    )

    pdfmetrics.registerFont(TTFont("CockpitCJK", font_path(), subfontIndex=0))
    pdfmetrics.registerFontFamily(
        "CockpitCJK",
        normal="CockpitCJK",
        bold="CockpitCJK",
        italic="CockpitCJK",
        boldItalic="CockpitCJK",
    )
    body = ParagraphStyle(
        "body",
        fontName="CockpitCJK",
        fontSize=10.5,
        leading=16,
        wordWrap="CJK",
        spaceAfter=6,
    )
    small = ParagraphStyle("cell", parent=body, fontSize=9, leading=12, spaceAfter=0)
    flow = []
    for kind, value in blocks(content):
        if kind == "table":
            table = LongTable(
                [[Paragraph(inline(c), small) for c in row] for row in value],
                colWidths=[(A4[0] - 96) / len(value[0])] * len(value[0]),
                repeatRows=1,
                hAlign="LEFT",
            )
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EAF0F8")),
                        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 6),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                    ]
                )
            )
            flow.extend([table, Spacer(1, 10)])
        else:
            style = body
            if kind.startswith("h"):
                size = 20 - min(int(kind[1]), 5) * 2
                style = ParagraphStyle(
                    kind,
                    parent=body,
                    fontSize=size,
                    leading=size * 1.5,
                    spaceBefore=10,
                    spaceAfter=8,
                )
            flow.append(Paragraph(inline(value), style))
    result = io.BytesIO()
    SimpleDocTemplate(
        result,
        pagesize=A4,
        topMargin=48,
        bottomMargin=48,
        leftMargin=48,
        rightMargin=48,
        title="",
    ).build(flow)
    return result.getvalue()


def word(content):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, Cm

    document = Document()
    for style_name in ("Normal", "Title", *(f"Heading {i}" for i in range(1, 7))):
        style = document.styles[style_name]
        style.font.name = "WenQuanYi Zen Hei"
        style.element.get_or_add_rPr().get_or_add_rFonts().set(
            qn("w:eastAsia"), "WenQuanYi Zen Hei"
        )
    document.styles["Normal"].font.size = Pt(11)
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    for kind, value in blocks(content):
        if kind == "table":
            table = document.add_table(rows=0, cols=len(value[0]), style="Table Grid")
            for row_index, row in enumerate(value):
                cells = table.add_row().cells
                for cell, text in zip(cells, row):
                    cell.text = text
                if row_index == 0:
                    repeat = OxmlElement("w:tblHeader")
                    table.rows[0]._tr.get_or_add_trPr().append(repeat)
        else:
            paragraph = (
                document.add_heading(level=int(kind[1]))
                if kind.startswith("h")
                else document.add_paragraph()
            )
            for index, part in enumerate(re.split(r"\*\*([^*]+)\*\*", value)):
                paragraph.add_run(part).bold = bool(index % 2)
    result = io.BytesIO()
    document.save(result)
    return result.getvalue()


def workbook(content):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    value = json.loads(content)
    fields(value, ("sheets",))
    if not isinstance(value["sheets"], list) or not 1 <= len(value["sheets"]) <= 20:
        raise ValueError("Invalid sheets")
    book = Workbook()
    book.remove(book.active)
    seen, cells = set(), 0
    for sheet in value["sheets"]:
        fields(sheet, ("name", "rows"))
        name = string(sheet["name"], 31)
        if (
            not name.strip()
            or re.search(r"[\\/*?:\[\]]", name)
            or name.casefold() in seen
        ):
            raise ValueError("Invalid sheet name")
        seen.add(name.casefold())
        rows = sheet["rows"]
        if not isinstance(rows, list) or len(rows) > 10000:
            raise ValueError("Invalid rows")
        ws = book.create_sheet(name)
        widths = {}
        for r, row in enumerate(rows, 1):
            if not isinstance(row, list) or len(row) > 100:
                raise ValueError("Invalid row")
            cells += len(row)
            if cells > 20000:
                raise ValueError("Too many cells")
            for c, value in enumerate(row, 1):
                cell = ws.cell(r, c)
                if isinstance(value, dict):
                    fields(value, ("formula",))
                    formula = string(value["formula"], 500)
                    # Explicit, local arithmetic/aggregate formulas only. No
                    # external workbooks, DDE, URLs, names or executable links.
                    if not re.fullmatch(r"=[A-Za-z0-9_\s.,():+*/^%=<>!\-]+", formula):
                        raise ValueError("Invalid formula")
                    functions = re.findall(r"([A-Za-z_][A-Za-z0-9_.]*)\s*\(", formula)
                    if any(
                        f.upper()
                        not in {
                            "SUM",
                            "AVERAGE",
                            "MIN",
                            "MAX",
                            "COUNT",
                            "COUNTA",
                            "ROUND",
                            "ROUNDUP",
                            "ROUNDDOWN",
                            "IF",
                            "ABS",
                        }
                        for f in functions
                    ):
                        raise ValueError("Unsupported formula")
                    # Every bare identifier must be a cell reference or one of
                    # the above functions; Excel computes only when opened.
                    remainder = re.sub(r"[A-Za-z_][A-Za-z0-9_.]*\s*\(", "(", formula)
                    remainder = re.sub(
                        r"\b[A-Za-z]{1,3}[1-9][0-9]{0,6}\b", "", remainder
                    )
                    if re.search(r"[A-Za-z_!]", remainder):
                        raise ValueError("Unsupported formula reference")
                    cell.value = formula
                elif value is None or type(value) is bool:
                    cell.value = value
                elif type(value) in (int, float):
                    cell.value = number(value, -1e100, 1e100)
                else:
                    cell.value = string(value, 10000)
                    # Even '=...', '@...', '+...' are literal strings unless
                    # the model explicitly used the validated formula object.
                    cell.data_type = "s"
                if r == 1:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="334155")
                widths[c] = min(45, max(widths.get(c, 12), len(str(value)) + 2))
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions
        for c, width in widths.items():
            ws.column_dimensions[get_column_letter(c)].width = width
    result = io.BytesIO()
    book.save(result)
    return result.getvalue()


def image(content, extension):
    from PIL import Image, ImageDraw, ImageFont

    value = json.loads(content)
    fields(value, ("width", "height", "elements"), ("background",))
    width, height = (
        number(value[k], 64, 4096, integer=True) for k in ("width", "height")
    )
    if (
        width * height > 8_000_000
        or not isinstance(value["elements"], list)
        or not 1 <= len(value["elements"]) <= 500
    ):
        raise ValueError("Canvas limits")

    def color(value):
        if not isinstance(value, str) or not re.fullmatch(r"#[0-9a-fA-F]{6}", value):
            raise ValueError("Expected #RRGGBB")
        return value

    canvas = Image.new(
        "RGB", (width, height), color(value.get("background", "#FFFFFF"))
    )
    draw = ImageDraw.Draw(canvas)
    fonts = {}
    for item in value["elements"]:
        if not isinstance(item, dict):
            raise ValueError("Invalid element")
        kind = item.get("type")
        if kind == "text":
            fields(item, ("type", "x", "y", "text"), ("size", "color"))
            point = (number(item["x"], 0, width), number(item["y"], 0, height))
            size = number(item.get("size", 28), 8, 160, integer=True)
            if size not in fonts:
                fonts[size] = ImageFont.truetype(font_path(), size)
            draw.multiline_text(
                point,
                string(item["text"]),
                font=fonts[size],
                fill=color(item.get("color", "#172033")),
                spacing=6,
            )
        elif kind in ("rect", "ellipse", "line"):
            fields(item, ("type", "box"), ("fill", "outline", "width"))
            box = item["box"]
            if not isinstance(box, list) or len(box) != 4:
                raise ValueError("Invalid box")
            box = [
                number(v, 0, width if i % 2 == 0 else height) for i, v in enumerate(box)
            ]
            stroke = number(item.get("width", 2), 1, 64, integer=True)
            outline = color(item.get("outline", "#334155"))
            if kind == "line":
                if "fill" in item:
                    raise ValueError("Line uses outline")
                draw.line(box, fill=outline, width=stroke)
            else:
                if box[2] < box[0] or box[3] < box[1]:
                    raise ValueError("Invalid box order")
                fill = color(item["fill"]) if "fill" in item else None
                if kind == "rect":
                    draw.rectangle(box, fill=fill, outline=outline, width=stroke)
                else:
                    draw.ellipse(box, fill=fill, outline=outline, width=stroke)
        else:
            raise ValueError("Unsupported element")
    result = io.BytesIO()
    fmt = {".png": "PNG", ".jpg": "JPEG", ".jpeg": "JPEG", ".webp": "WEBP"}[extension]
    canvas.save(result, format=fmt, **({"quality": 92} if fmt != "PNG" else {}))
    return result.getvalue()


def main():
    resource.setrlimit(resource.RLIMIT_CPU, (12, 13))
    resource.setrlimit(resource.RLIMIT_FSIZE, (MAX_BYTES, MAX_BYTES))
    resource.setrlimit(resource.RLIMIT_CORE, (0, 0))
    if sys.platform != "darwin":
        resource.setrlimit(resource.RLIMIT_AS, (512 * 1024 * 1024,) * 2)
    raw = sys.stdin.buffer.read(131073)
    if len(raw) > 131072:
        raise ValueError("Input limit")
    value = json.loads(raw)
    fields(value, ("extension", "content"))
    content = string(value["content"], 65536)
    ext = value["extension"]
    if ext == ".pdf":
        result = pdf(content)
    elif ext == ".docx":
        result = word(content)
    elif ext == ".xlsx":
        result = workbook(content)
    elif ext in {".png", ".jpg", ".jpeg", ".webp"}:
        result = image(content, ext)
    else:
        raise ValueError("Unsupported format")
    if not 0 < len(result) <= MAX_BYTES:
        raise ValueError("Output limit")
    sys.stdout.buffer.write(result)


if __name__ == "__main__":
    try:
        main()
    except Exception:
        # No model content, host path or traceback in tool errors.
        sys.exit(2)
