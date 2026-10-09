"""Selected audit repairs only; no remote quota/notification calls."""

from session_fixture import fixture_identity, session_auth

import json
import os
import unittest
import uuid
from copy import deepcopy
from datetime import datetime, date, timedelta
from decimal import Decimal
from unittest.mock import patch

import psycopg
from psycopg import sql
from psycopg.rows import dict_row
from openpyxl import load_workbook

from new_api_cockpit import balance, scopes, quota, quota_schedule, report_snapshots
from new_api_cockpit.app import app
from new_api_cockpit import metadata_fallback as fallback
from new_api_cockpit import runtime
from new_api_cockpit import report


def log(index=1, **changes):
    return (
        dict(
            id=index,
            created_at=index,
            user_id=1,
            username="fixture",
            token_id=10,
            token_name="fixture-key",
            model_name="gpt",
            group_name="auto",
            channel_id=1,
            quota=500000,
            prompt_tokens=100,
            completion_tokens=20,
            other=json.dumps(
                dict(
                    cache_tokens=30,
                    cache_write_tokens=10,
                    group_ratio=3,
                    model_ratio=1,
                    completion_ratio=5,
                    cache_ratio="0.1",
                    cache_creation_ratio="1.25",
                )
            ),
        )
        | changes
    )


class RepairUnitTest(unittest.TestCase):
    def test_pathological_amounts_are_json_400_before_any_network(self):
        for value in (
            "sNaN",
            "NaN",
            "Infinity",
            "-Infinity",
            "1e9999999",
            "1e-9999999",
            "9" * 129,
            "1.0000000000000000000000000000000000000001",
            "1e10000000000000000000000",
        ):
            with (
                self.subTest(value=value),
                patch(
                    "new_api_cockpit.app.request_identity", side_effect=fixture_identity
                ),
                patch.object(quota, "connect") as connect,
            ):
                response = app.test_client().post(
                    "/cockpit/api/users/quota/preview",
                    auth=session_auth("admin"),
                    headers={"X-Quota-Action": "preview"},
                    json=dict(user_ids=[1], mode="add", amount_yuan=value),
                )
                self.assertEqual(response.status_code, 400)
                self.assertTrue(response.is_json)
                connect.assert_not_called()
        self.assertEqual(
            quota.validate_request(
                dict(user_ids=[1], mode="add", amount_yuan="1.123456")
            )[2],
            561728,
        )
        self.assertEqual(
            quota.validate_request(
                dict(user_ids=[1], mode="add", amount_yuan="1000000000")
            )[2],
            500000000000000,
        )

    def test_every_external_excel_string_is_literal_and_generated_formulas_remain(self):
        row = report.decorate(
            fallback.aggregate([log(model_name="=1+1", username="=2+2")], False), {}
        )[0]
        row["display_name"] = "=3+3"
        row["tier_name"] = "=4+4"
        row["pricing_mode"] = "expression"
        row["price_tiers"] = [
            dict(
                name="=5+5",
                condition="=6+6",
                input_price=Decimal(2),
                output_price=Decimal(10),
                cache_price=Decimal(".2"),
                write_price=Decimal("2.5"),
            )
        ]
        workbook = load_workbook(report.export_excel([row], "2026-09-01", "2026-09-01"))
        count = 0
        for sheet in workbook:
            for cells in sheet:
                for cell in cells:
                    if cell.value in ("=1+1", "=2+2", "=3+3", "=4+4", "=5+5", "=6+6"):
                        self.assertEqual(
                            cell.data_type, "s", (sheet.title, cell.coordinate)
                        )
                        count += 1
        self.assertGreaterEqual(count, 10)
        self.assertEqual(workbook["用户模型用量"]["P4"].data_type, "f")
        self.assertEqual(workbook["区间汇总"]["B3"].data_type, "f")
        self.assertEqual(workbook["表达式价格"]["D3"].data_type, "n")
        self.assertEqual(workbook["模型消费"]["B4"].value, None)

    def test_corrupt_metadata_keeps_fee_and_unknown_counts_are_not_zero(self):
        records = [
            log(1),
            log(2, other="{broken"),
            log(3, model_name="claude-test", other="[]"),
            log(4, other='{"cache_tokens":"oops"}'),
            log(5, other='{"group_ratio":"NaN"}'),
        ]
        rows = report.decorate(fallback.aggregate(records, False), {})
        total = report.totals(rows, "2026-09-01", "2026-09-01")
        self.assertEqual(total["amount"], 5)
        self.assertEqual(total["request_count"], 5)
        self.assertIsNone(total["total_tokens"])
        self.assertIsNone(total["cache_read_tokens"])
        self.assertEqual(total["output_tokens"], 100)
        self.assertIsNone(total["tpm"])
        unknown = [r for r in rows if r.get("metadata_error_count")]
        self.assertEqual(sum(r["request_count"] for r in unknown), 4)
        self.assertTrue(all(r["cost_formula"]["calculated"] is None for r in unknown))
        self.assertTrue(all(r["input_price"] is None for r in unknown))
        ranked = report.rankings(rows)
        self.assertIsNone(ranked["user_tokens"][0]["total_tokens"])
        wb = load_workbook(report.export_excel(rows, "2026-09-01", "2026-09-01"))
        self.assertIsNone(wb["用户模型用量"][f"F{len(rows) + 3}"].value)
        self.assertEqual(wb["用户模型用量"][f"P{len(rows) + 3}"].data_type, "f")
        self.assertIsNone(wb["区间汇总"]["B4"].value)
        self.assertIsNone(wb["区间汇总"]["B11"].value)
        for broken in ("oops", "[1]", '{"status_code":null}'):
            failures = fallback.aggregate([log(other=broken)], False, failures=True)
            self.assertEqual(failures[0]["status_code"], "未知")
            self.assertEqual(failures[0]["failure_count"], 1)

    def test_runtime_migrates_before_exec_and_stops_on_failure(self):
        events = []
        with (
            patch.object(balance, "configured", return_value=True),
            patch.object(
                balance, "initialize", side_effect=lambda: events.append("migrate")
            ),
            patch.object(
                runtime.os, "execvp", side_effect=lambda *_: events.append("exec")
            ),
            patch.object(runtime.sys, "argv", ["runtime", "gunicorn", "module:app"]),
        ):
            runtime.main()
        self.assertEqual(events, ["migrate", "exec"])
        with (
            patch.object(balance, "configured", return_value=True),
            patch.object(
                balance, "initialize", side_effect=RuntimeError("not for logs")
            ),
            patch.object(runtime.os, "execvp") as execute,
            self.assertLogs(level="ERROR") as captured,
            self.assertRaises(SystemExit) as error,
        ):
            runtime.main()
        self.assertEqual(error.exception.code, 1)
        execute.assert_not_called()
        self.assertNotIn("not for logs", "".join(captured.output))
        with (
            patch.object(balance, "configured", return_value=False),
            patch.object(balance, "initialize") as migrate,
            patch.object(runtime.os, "execvp") as execute,
            patch.object(runtime.sys, "argv", ["runtime", "gunicorn", "module:app"]),
        ):
            runtime.main()
        migrate.assert_not_called()
        execute.assert_called_once_with(
            "gunicorn",
            [
                "gunicorn",
                "--config",
                "python:new_api_cockpit.gunicorn_conf",
                "module:app",
            ],
        )

    def test_schema_check_never_invokes_migration_on_rule_reads(self):
        with (
            patch.object(balance, "configured", return_value=True),
            patch.object(balance, "require_schema") as check,
            patch.object(balance, "initialize") as migrate,
        ):
            quota_schedule.initialize()
            check.assert_called_once()
            migrate.assert_not_called()


@unittest.skipUnless(os.environ.get("MONITOR_DATABASE_URL"), "disposable PG required")
class PerformanceDatabaseTest(unittest.TestCase):
    def setUp(self):
        self.dsn = os.environ["MONITOR_DATABASE_URL"]
        self.schema = "repairs_" + uuid.uuid4().hex
        with psycopg.connect(self.dsn) as conn:
            conn.execute(
                sql.SQL("CREATE SCHEMA {}").format(sql.Identifier(self.schema))
            )
        self.addCleanup(self.cleanup)
        self.enterContext(patch.object(balance, "connect", self.connect))
        self.enterContext(patch.object(balance, "configured", return_value=True))
        balance.initialize()
        self.catalog = [
            dict(channel_id=1, channel_name="one", channel_status=1, tag_value="A"),
            dict(channel_id=2, channel_name="two", channel_status=2, tag_value="B"),
        ]
        self.source = self.enterContext(
            patch.object(balance, "source_channels", return_value=self.catalog)
        )
        self.enterContext(patch("new_api_cockpit.notifications.notify_safely"))

    def connect(self):
        return psycopg.connect(
            self.dsn, row_factory=dict_row, options="-c search_path=" + self.schema
        )

    def cleanup(self):
        with psycopg.connect(self.dsn) as conn:
            conn.execute(
                sql.SQL("DROP SCHEMA {} CASCADE").format(sql.Identifier(self.schema))
            )

    def test_catalog_discovers_history_once_and_repeated_reads_do_not_update_rows(self):
        scopes.refresh_scopes()
        self.assertEqual(self.source.call_args.kwargs, {"include_deleted": True})
        with self.connect() as conn:
            before = conn.execute(
                "SELECT channel_id,xmin::text AS xmin,last_seen_at FROM balance_channel_inventory ORDER BY channel_id"
            ).fetchall()
            sequence_before = conn.execute(
                "SELECT last_value FROM balance_scopes_id_seq"
            ).fetchone()
            scope_before = conn.execute(
                "SELECT id,xmin::text AS xmin,updated_at FROM balance_scopes ORDER BY id"
            ).fetchall()
        scopes.refresh_scopes()
        self.assertEqual(self.source.call_args.kwargs, {"include_deleted": False})
        with self.connect() as conn:
            self.assertEqual(
                before,
                conn.execute(
                    "SELECT channel_id,xmin::text AS xmin,last_seen_at FROM balance_channel_inventory ORDER BY channel_id"
                ).fetchall(),
            )
            self.assertEqual(
                scope_before,
                conn.execute(
                    "SELECT id,xmin::text AS xmin,updated_at FROM balance_scopes ORDER BY id"
                ).fetchall(),
            )
            self.assertEqual(
                sequence_before,
                conn.execute("SELECT last_value FROM balance_scopes_id_seq").fetchone(),
            )
        scopes.refresh_scopes(full_discovery=True)
        self.assertEqual(self.source.call_args.kwargs, {"include_deleted": True})
        self.source.reset_mock()
        balance.usage_channels_snapshot()
        self.assertEqual(self.source.call_count, 1)

    def test_catalog_source_failure_does_not_mark_history_as_discovered(self):
        self.source.side_effect = RuntimeError("fixture source unavailable")
        with self.assertRaises(RuntimeError):
            scopes.refresh_scopes()
        with self.connect() as conn:
            self.assertFalse(
                conn.execute(
                    "SELECT history_discovered FROM channel_catalog_sync_state WHERE id=1"
                ).fetchone()["history_discovered"]
            )
            self.assertEqual(
                conn.execute(
                    "SELECT count(*) AS n FROM balance_channel_inventory"
                ).fetchone()["n"],
                0,
            )
        self.source.side_effect = None
        scopes.refresh_scopes()
        self.assertEqual(self.source.call_args.kwargs, {"include_deleted": True})

    def test_daily_all_ledgers_share_one_query_and_manual_is_fresh(self):
        now = datetime(2026, 9, 20, 10, tzinfo=report.TZ)
        current = date(2026, 9, 1)
        scopes.refresh_scopes()
        with self.connect() as conn:
            conn.execute(
                "UPDATE balance_settings SET enabled=true,start_month=%s,budget=1000",
                (current,),
            )
        details = {
            current: [
                dict(channel_id=1, amount=Decimal(10)),
                dict(channel_id=2, amount=Decimal(20)),
                dict(channel_id=99, amount=Decimal(5)),
            ]
        }
        with patch.object(
            balance, "source_channel_amounts", return_value=details
        ) as query:
            result = balance.check_all_enabled(now)
            self.assertTrue(all(result.values()))
            query.assert_called_once_with([current], now)
            self.assertTrue(
                all(value is None for value in balance.check_all_enabled(now).values())
            )
            self.assertEqual(query.call_count, 1)
            tag = scopes.list_scopes(False)[1]["id"]
            balance.check_once(now, scope_id=tag, daily=False)
            self.assertEqual(query.call_count, 2)
        with self.connect() as conn:
            rows = conn.execute(
                "SELECT s.kind,s.tag_value,b.current_amount FROM balance_scopes s JOIN balance_state b ON b.scope_id=s.id"
            ).fetchall()
        self.assertEqual(
            {r["tag_value"] or r["kind"]: r["current_amount"] for r in rows},
            {"all": 35, "A": 10, "B": 20, "ungrouped": 5},
        )

    def test_lazy_details_are_immutable_scoped_owned_and_expire(self):
        original = report.decorate(fallback.aggregate([log()], False), {})
        slim = report_snapshots.create(original, "alice", 1)
        self.assertTrue(report_snapshots.HEAVY_FIELDS.isdisjoint(slim[0]))
        body = dict(report_id=slim[0]["report_id"], row_ids=[0])
        before = deepcopy(original[0]["cost_formula"])
        original[0]["cost_formula"]["buckets"][0]["actual"] = "changed"
        detail = report_snapshots.fetch(body, "alice", 1)[0]["detail"]["cost_formula"]
        self.assertEqual(detail, json.loads(report_snapshots.dumps(before)))
        for owner, scope_id in [("bob", 1), ("alice", 2)]:
            with self.assertRaises(report_snapshots.SnapshotExpired):
                report_snapshots.fetch(body, owner, scope_id)
        with self.connect() as conn:
            conn.execute(
                "UPDATE report_snapshots SET expires_at=%s",
                (datetime.now(report.TZ) - timedelta(seconds=1),),
            )
        with self.assertRaises(report_snapshots.SnapshotExpired):
            report_snapshots.fetch(body, "alice", 1)
        for _ in range(report_snapshots.MAX_REPORTS_PER_OWNER + 2):
            report_snapshots.create(original, "alice", 1)
        with self.connect() as conn:
            self.assertEqual(
                conn.execute("SELECT count(*) AS n FROM report_snapshots").fetchone()[
                    "n"
                ],
                report_snapshots.MAX_REPORTS_PER_OWNER,
            )

    def test_lazy_api_keeps_default_response_and_enforces_detail_access(self):
        rows = report.decorate(fallback.aggregate([log()], False), {})
        client = app.test_client()
        with (
            patch("new_api_cockpit.app.request_identity", side_effect=fixture_identity),
            patch("new_api_cockpit.app.load_report", return_value=rows),
        ):
            url = "/cockpit/api/statistics/usage?start=2026-09-01&end=2026-09-01"
            normal = client.get(url, auth=session_auth("alice"))
            self.assertIn("cost_formula", normal.json["rows"][0])
            lazy = client.get(url + "&details=lazy", auth=session_auth("alice"))
            self.assertEqual(lazy.status_code, 200)
            row = lazy.json["rows"][0]
            self.assertNotIn("cost_formula", row)
            self.assertEqual(lazy.json["totals"], normal.json["totals"])
            self.assertEqual(lazy.json["rankings"], normal.json["rankings"])
            payload = dict(report_id=row["report_id"], row_ids=[row["row_id"]])
            route = "/cockpit/api/statistics/usage/details"
            self.assertEqual(
                client.post(
                    route, json=payload, auth=session_auth("alice")
                ).status_code,
                403,
            )
            for username, expected in [("alice", 200), ("bob", 410)]:
                response = client.post(
                    route,
                    json=payload,
                    auth=session_auth(username),
                    headers={"X-Statistics-Request": "1"},
                )
                self.assertEqual(response.status_code, expected)
            for ids in ([True], [-1], [2147483648], [], [0] * 501):
                response = client.post(
                    route,
                    json=dict(payload, row_ids=ids),
                    auth=session_auth("alice"),
                    headers={"X-Statistics-Request": "1"},
                )
                self.assertEqual(response.status_code, 400)
            cross = client.post(
                route,
                json=payload,
                auth=session_auth("alice"),
                headers={"X-Statistics-Request": "1", "Sec-Fetch-Site": "cross-site"},
            )
            self.assertEqual(cross.status_code, 403)

    def test_require_schema_reads_only_and_fails_closed_without_ddl(self):
        balance.require_schema()
        with self.connect() as conn:
            conn.execute("DELETE FROM schema_migrations WHERE version LIKE '008%%'")
        with (
            patch.object(balance, "initialize") as migrate,
            self.assertRaises(quota_schedule.ScheduleUnavailable),
        ):
            quota_schedule.initialize()
        migrate.assert_not_called()

    def test_real_fallback_matches_healthy_sql_and_handles_broken_json_same_snapshot(
        self,
    ):
        # Disposable source fixture, not the live source database.
        records = [
            log(1),
            log(2, other=json.dumps(dict(cache_tokens=5, group_ratio="3.4"))),
        ]
        with self.connect() as conn:
            conn.execute("""CREATE TABLE logs(id bigint,created_at bigint,user_id bigint,username text,
                token_id bigint,token_name text,model_name text,"group" text,channel_id bigint,
                quota bigint,prompt_tokens bigint,completion_tokens bigint,other text,type integer);
                CREATE TABLE users(id bigint,display_name text); CREATE TABLE options(key text,value text)""")
            for r in records:
                conn.execute(
                    "INSERT INTO logs VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,2)",
                    [
                        r[k]
                        for k in (
                            "id",
                            "created_at",
                            "user_id",
                            "username",
                            "token_id",
                            "token_name",
                            "model_name",
                            "group_name",
                            "channel_id",
                            "quota",
                            "prompt_tokens",
                            "completion_tokens",
                            "other",
                        )
                    ],
                )
            params = dict(
                start=0,
                end=100,
                by_token=True,
                token_ids=None,
                groups=None,
                channel_ids=None,
                excluded_channel_ids=None,
            )
            baseline = conn.execute(report.SQL, params).fetchall()
            for row in baseline:
                row.pop("metadata_invalid", None)
            expected = report.decorate(baseline, {})
            actual = report.decorate(fallback.aggregate(records, True), {})
            # Nested snapshots are deliberately text in SQL; numeric comparisons
            # focus on all displayed totals, grouping and formula amounts.
            fields = [
                "username",
                "model_name",
                "token_id",
                "token_name",
                "request_count",
                "ratio_count",
                "amount",
                "total_tokens",
                *report.TOKEN_FIELDS,
            ]
            self.assertEqual(
                [{k: r[k] for k in fields} for r in actual],
                [{k: r[k] for k in fields} for r in expected],
            )
            conn.execute(
                "INSERT INTO logs SELECT 3,3,1,'fixture',10,'fixture-key','gpt','auto',1,500000,100,20,'broken',2"
            )
            conn.execute(
                "INSERT INTO logs SELECT 4,4,1,'fixture',10,'fixture-key','gpt','auto',1,0,0,0,'broken',5"
            )

            conn.execute(
                "INSERT INTO logs SELECT 5,5,2,'unselected',99,'other-key','gpt','other',2,5000000,100,20,'broken',2"
            )

        def source_connection(**kwargs):
            self.assertIn("default_transaction_read_only=on", kwargs["options"])
            return psycopg.Connection.connect(
                self.dsn,
                row_factory=dict_row,
                options="-c search_path="
                + self.schema
                + " -c default_transaction_read_only=on",
            )

        with patch.object(report.psycopg, "connect", side_effect=source_connection):
            rows = report.load_report(
                "1970-01-01T08:00:00",
                "1970-01-01T08:01:00",
                by_token=True,
                include_failures=True,
                channel_ids=[1],
                token_ids=[10],
                groups=["auto"],
            )
        self.assertEqual(sum(r["amount"] for r in rows), 3)
        self.assertEqual(sum(r["request_count"] for r in rows), 3)
        self.assertEqual(sum(r["failure_count"] for r in rows), 1)
        self.assertEqual(sum(r.get("metadata_error_count", 0) for r in rows), 1)
