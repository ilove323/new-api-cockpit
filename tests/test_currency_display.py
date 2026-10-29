"""Read-only exchange metadata and RMB-only numeric exports; no upstream API calls."""

from copy import deepcopy
from decimal import Decimal
from unittest import TestCase
from unittest.mock import MagicMock, patch

from openpyxl import load_workbook

from new_api_cockpit import currency_display, report
from new_api_cockpit.app import app
from session_fixture import fixture_identity, session_auth

ROW = dict(
    user_id=1,
    username="fixture",
    display_name="测试",
    request_count=2,
    model_name="model-a",
    tier_name="low",
    total_tokens=1000,
    input_tokens=800,
    output_tokens=100,
    cache_read_tokens=100,
    cache_write_tokens=0,
    group_ratio=Decimal("1.5"),
    input_price=Decimal("2"),
    output_price=Decimal("10"),
    cache_price=Decimal("0.2"),
    write_price=None,
    amount=Decimal("1.25"),
    pricing_mode="expression",
    price_tiers=[
        dict(
            name="low",
            condition="",
            input_price=Decimal("2"),
            output_price=Decimal("10"),
            cache_price=Decimal("0.2"),
            write_price=None,
        )
    ],
)


class CurrencyDisplayTest(TestCase):
    def test_configured_rate_beats_default_and_payment_price_is_not_an_exchange_rate(
        self,
    ):
        self.assertEqual(currency_display.settings({})["usd_cny_rate"], Decimal("7.3"))
        self.assertEqual(
            currency_display.settings({"USDExchangeRate": "7.21", "Price": "5"})[
                "usd_cny_rate"
            ],
            Decimal("7.21"),
        )
        self.assertEqual(
            currency_display.settings({"Price": "5"})["rate_source"], "new_api_default"
        )
        for value in (
            "",
            None,
            "invalid",
            "NaN",
            "Infinity",
            "0",
            "-1",
            "1000001",
            "1e9999999",
            "1e-9999999",
        ):
            with self.subTest(value=value):
                self.assertIsNone(
                    currency_display.settings({"USDExchangeRate": value})[
                        "usd_cny_rate"
                    ]
                )
                with self.assertRaises(ValueError):
                    report.export_excel(
                        [ROW], "2026-10-01", "2026-10-02", usd_cny_rate=value
                    )

    def test_rate_reads_share_the_existing_options_query_and_repeatable_read_connection(
        self,
    ):
        connection = MagicMock()
        conn = connection.__enter__.return_value
        conn.execute.return_value.fetchall.side_effect = [
            [],
            [
                {"key": "ModelRatio", "value": "{}"},
                {"key": "USDExchangeRate", "value": "7.2"},
            ],
        ]
        metadata = {}
        with patch.object(
            report.psycopg, "connect", return_value=connection
        ) as connect:
            self.assertEqual(
                report.load_report(
                    "2026-10-01", "2026-10-02", currency_display=metadata
                ),
                [],
            )
        connect.assert_called_once()
        calls = conn.execute.call_args_list
        self.assertEqual(len(calls), 3)  # transaction mode, aggregate, options
        self.assertIn("USDExchangeRate", calls[-1].args[1][0])
        self.assertEqual(metadata["usd_cny_rate"], Decimal("7.2"))
        self.assertIn(
            "default_transaction_read_only=on", connect.call_args.kwargs["options"]
        )

    def test_all_excel_money_is_rmb_and_original_usage_prices_and_ratios_are_untouched(
        self,
    ):
        rows = [deepcopy(ROW)]
        before = deepcopy(rows)
        wb = load_workbook(
            report.export_excel(
                rows, "2026-10-01", "2026-10-02", usd_cny_rate=Decimal("7.2")
            )
        )
        sheet = wb["用户模型用量"]
        self.assertEqual(sheet["L3"].value, 14.4)
        self.assertEqual(sheet["M3"].value, 72)
        self.assertEqual(sheet["N3"].value, 1.44)
        self.assertIsNone(sheet["O3"].value)
        self.assertEqual(sheet["P3"].value, 9)
        self.assertEqual(sheet["P3"].data_type, "n")
        self.assertEqual(sheet["K3"].value, 1.5)
        self.assertEqual(sheet["F3"].value, 1000)
        self.assertIn("RMB", sheet["C1"].value)
        self.assertIn("7.2", sheet["C1"].value)
        self.assertEqual(wb["模型消费"]["C3"].value, 9)
        self.assertEqual(wb["用户消费"]["D3"].value, 9)
        self.assertEqual(wb["表达式价格"]["D3"].value, 14.4)
        self.assertEqual(sheet["P4"].value, "=SUM(P3:P3)")
        self.assertEqual(wb["区间汇总"]["B3"].value, "='用户模型用量'!P4")
        self.assertEqual(rows, before)

    def test_http_usage_keeps_raw_values_and_exports_ignore_client_currency(self):
        rows = [deepcopy(ROW)]

        def load(start, end, **kwargs):
            if kwargs.get("currency_display") is not None:
                kwargs["currency_display"].update(
                    currency_display.settings({"USDExchangeRate": "7.2"})
                )
            return rows

        args = dict(auth=session_auth("fixture_admin"))
        query = "?start=2026-10-01&end=2026-10-02&currency=USD"
        with (
            patch("new_api_cockpit.app.request_identity", side_effect=fixture_identity),
            patch("new_api_cockpit.app.balance.configured", return_value=False),
            patch("new_api_cockpit.app.load_report", side_effect=load) as loader,
        ):
            client = app.test_client()
            usage = client.get("/cockpit/api/statistics/usage" + query, **args)
            self.assertEqual(usage.status_code, 200)
            self.assertEqual(usage.json["currency_display"]["usd_cny_rate"], "7.2")
            self.assertEqual(usage.json["currency_display"]["default_currency"], "CNY")
            self.assertEqual(usage.json["rows"][0]["amount"], "1.25")
            self.assertEqual(usage.json["rows"][0]["input_price"], "2")
            self.assertEqual(usage.json["totals"]["amount"], "1.25")
            loader.reset_mock()
            result = client.get("/cockpit/api/statistics/export" + query, **args)
            self.assertEqual(result.status_code, 200)
            from io import BytesIO

            wb = load_workbook(BytesIO(result.data))
            self.assertEqual(wb.active["P3"].value, 9)
            loader.assert_called_once()
            self.assertEqual(rows[0]["amount"], Decimal("1.25"))
            result.close()
