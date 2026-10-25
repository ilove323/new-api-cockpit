"""Real binary generation/reopen + complete synthetic-model tool lifecycle."""

from io import BytesIO
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from uuid import uuid4
import zipfile

from docx import Document
from openpyxl import load_workbook
from PIL import Image
import pypdfium2 as pdfium

from chat_fixtures import access, body, create, make_service
from test_chat_storage import StorageFixture
from new_api_cockpit.features.chat import config, media, outputs, tools
from new_api_cockpit.features.chat.errors import ChatError
from new_api_cockpit.features.chat.storage import ChatStore
from new_api_cockpit.features.chat.workspace import Workspace

TEXT = "# 核验报告\n中文与 English 123\n\n**结论**：真正的标准文件。\n\n|项目|金额|\n|---|---|\n|研发|42|"
BOOK = json.dumps(
    {
        "sheets": [
            {
                "name": "预算",
                "rows": [
                    ["项目", "金额"],
                    ["研发", 42],
                    ["合计", {"formula": "=SUM(B2:B2)"}],
                    ['=WEBSERVICE("https://invalid")', None],
                ],
            }
        ]
    },
    ensure_ascii=False,
)
DRAW = json.dumps(
    {
        "width": 640,
        "height": 360,
        "background": "#FFFFFF",
        "elements": [
            {"type": "text", "x": 24, "y": 20, "text": "预算示意图 · 中文", "size": 28},
            {"type": "rect", "box": [50, 90, 180, 300], "fill": "#2563EB"},
            {"type": "ellipse", "box": [240, 100, 400, 260], "fill": "#10B981"},
            {"type": "line", "box": [40, 310, 590, 310], "outline": "#333333"},
        ],
    },
    ensure_ascii=False,
)
SPECS = {
    "report.pdf": TEXT,
    "report.docx": TEXT,
    "report.xlsx": BOOK,
    "chart.png": DRAW,
}


def validate(test, name, raw):
    test.assertEqual(ChatStore.normalize_file(name, raw, "output"), raw)
    test.assertEqual(
        media.mime_for(name, output=True), config.BINARY_MIMES[Path(name).suffix]
    )
    if name.endswith(".pdf"):
        test.assertTrue(raw.startswith(b"%PDF-"))
        with pdfium.PdfDocument(raw) as doc:
            test.assertGreater(len(doc), 0)
            page = doc[0]
            text = page.get_textpage()
            try:
                test.assertIn("中文", text.get_text_range())
            finally:
                text.close()
                page.close()
        # TrueType font is embedded, not dependent on viewer's CJK fonts.
        test.assertIn(b"/FontFile2", raw)
    elif name.endswith(".docx"):
        test.assertTrue(zipfile.is_zipfile(BytesIO(raw)))
        doc = Document(BytesIO(raw))
        test.assertTrue(any("中文" in p.text for p in doc.paragraphs))
        test.assertEqual(doc.tables[0].cell(1, 1).text, "42")
    elif name.endswith(".xlsx"):
        with zipfile.ZipFile(BytesIO(raw)) as z:
            test.assertIn("xl/workbook.xml", z.namelist())
            test.assertFalse(
                any("externalLink" in p or "vbaProject" in p for p in z.namelist())
            )
        book = load_workbook(BytesIO(raw), data_only=False)
        try:
            ws = book["预算"]
            test.assertEqual(ws["B2"].value, 42)
            test.assertEqual(ws["B3"].data_type, "f")
            test.assertEqual(ws["B3"].value, "=SUM(B2:B2)")
            test.assertEqual(ws["A4"].data_type, "s")
        finally:
            book.close()
    else:
        with Image.open(BytesIO(raw)) as img:
            img.load()
            test.assertEqual(img.size, (640, 360))
            test.assertEqual(img.getpixel((80, 120)), (37, 99, 235))


class OutputRendererTest(unittest.TestCase):
    def test_real_files_open_in_independent_readers(self):
        for name, spec in SPECS.items():
            with self.subTest(name=name):
                validate(self, name, outputs.generate(name, spec))

    def test_image_codecs_are_real_not_png_with_different_name(self):
        for extension, fmt in [("jpg", "JPEG"), ("jpeg", "JPEG"), ("webp", "WEBP")]:
            raw = outputs.generate("chart." + extension, DRAW)
            with Image.open(BytesIO(raw)) as img:
                self.assertEqual(img.format, fmt)
                self.assertEqual(img.size, (640, 360))
            self.assertEqual(
                ChatStore.normalize_file("chart." + extension, raw, "output"), raw
            )

    def test_plain_text_named_binary_is_rejected_even_with_fake_magic(self):
        for name in SPECS:
            with self.subTest(name=name), self.assertRaises(ChatError):
                ChatStore.normalize_file(
                    name, b"%PDF- fake PK\x03\x04 not a file", "output"
                )
        for extension in ("doc", "xls", "gif", "bmp", "exe", "zip"):
            with self.assertRaises(ChatError):
                outputs.generate("bad." + extension, TEXT)

    def test_invalid_specs_and_code_urls_paths_cannot_execute(self):
        invalid = [
            ("x.xlsx", "{}"),
            ("x.png", "# Not a drawing spec"),
            ("x.png", json.dumps({"width": 9000, "height": 100, "elements": []})),
            ("x.png", json.dumps({"width": 4000, "height": 4000, "elements": []})),
            (
                "x.png",
                json.dumps(
                    {
                        "width": 640,
                        "height": 360,
                        "elements": [{"type": "image", "url": "file:///etc/passwd"}],
                    }
                ),
            ),
            ("x.xlsx", json.dumps({"sheets": [{"name": "bad/name", "rows": []}]})),
            (
                "x.xlsx",
                json.dumps(
                    {
                        "sheets": [
                            {"name": "A", "rows": [[{"formula": "=WEBSERVICE(A1)"}]]}
                        ]
                    }
                ),
            ),
            (
                "x.xlsx",
                json.dumps(
                    {
                        "sheets": [
                            {
                                "name": "A",
                                "rows": [[{"formula": "='[other.xlsx]Sheet1'!A1"}]],
                            }
                        ]
                    }
                ),
            ),
            (
                "x.xlsx",
                json.dumps(
                    {"sheets": [{"name": "A", "rows": [[{"formula": "=cmd|A1"}]]}]}
                ),
            ),
            (
                "x.xlsx",
                json.dumps(
                    {"sheets": [{"name": "A", "rows": [[{"formula": "=SECRET+1"}]]}]}
                ),
            ),
            (
                "x.xlsx",
                json.dumps(
                    {"sheets": [{"name": "A", "rows": []}, {"name": "a", "rows": []}]}
                ),
            ),
        ]
        for name, spec in invalid:
            with self.subTest(spec=spec), self.assertRaises(ChatError) as error:
                outputs.generate(name, spec)
            self.assertEqual(error.exception.code, "TOOL_OUTPUT_INVALID")

    def test_pdf_model_markup_is_literal_not_external_resource(self):
        raw = outputs.generate(
            "x.pdf",
            '# 中文\n<img src="file:///etc/passwd"/>\n<a href="https://example.invalid">hello</a>',
        )
        text = media.extract("x.pdf", raw)["text"]
        self.assertIn("file:///etc/passwd", text)
        self.assertNotIn("root:x:", text)

    def test_output_cap_and_timeout_fail_closed(self):
        with patch.object(config, "OUTPUT_FILE_BYTES", 1), self.assertRaises(ChatError):
            outputs.generate("x.png", DRAW)
        with (
            patch.object(config, "OUTPUT_RENDER_SECONDS", 0.000001),
            self.assertRaises(ChatError),
        ):
            outputs.generate("x.pdf", TEXT)


class OutputWorkspaceTest(StorageFixture):
    def test_versions_dedup_restart_download_and_readback(self):
        original = self.store.upload(
            self.scope, "report.pdf", outputs.generate("report.pdf", TEXT), str(uuid4())
        )
        with self.store.lease(self.scope, run=True) as lease:
            rid, _, _ = self.accept(lease)
            ws = Workspace(self.store, self.scope, lease, rid)
            results = []
            for i, (name, spec) in enumerate(SPECS.items()):
                result = tools.execute(
                    "write_file",
                    {"name": name, "content": spec},
                    ws,
                    {"run_id": rid, "call_id": f"out-{i}"},
                )
                self.assertTrue(result["ok"], result)
                results.append(result)
                entry, raw = self.store.read_file(
                    self.scope, result["file_id"], lease=lease
                )
                validate(self, name, raw)
                self.assertEqual(entry["kind"], "output")
                with patch.object(
                    outputs,
                    "generate",
                    side_effect=AssertionError("must not render twice"),
                ):
                    self.assertEqual(ws.write_file(name, spec, f"out-{i}"), result)
                self.assertEqual(
                    ws.write_file(name, spec + "changed", f"out-{i}")["code"],
                    "CHAT_IDEMPOTENCY_CONFLICT",
                )
            second = ws.write_file("report.pdf", TEXT + "\n第二版", "v2")
            self.assertEqual(second["version"], 2)
            self.assertEqual(original["kind"], "upload")
            self.terminal(lease, rid)
        self.store.close()
        self.store = ChatStore(self.settings)
        self.addCleanup(self.store.close)
        self.scope = self.store.open_conversation(self.a, self.cid)
        for result in results:
            entry, raw = self.store.read_file(self.scope, result["file_id"])
            validate(self, entry["name"], raw)
        self.assertIn(
            "中文",
            Workspace(self.store, self.scope).read_file(results[0]["file_id"])["text"],
        )
        self.assertEqual(len(self.store.files(self.scope)), 6)

    def test_failed_generation_never_publishes_output_or_file_added(self):
        with self.store.lease(self.scope, run=True) as lease:
            rid, _, _ = self.accept(lease)
            result = tools.execute(
                "write_file",
                {"name": "bad.png", "content": "not a drawing"},
                Workspace(self.store, self.scope, lease, rid),
                {"run_id": rid, "call_id": "bad"},
            )
            self.assertFalse(result["ok"])
            self.assertEqual(result["code"], "TOOL_OUTPUT_INVALID")
            self.assertEqual(self.store.files(self.scope, lease=lease), [])
            self.terminal(lease, rid)


class OutputRunTest(unittest.TestCase):
    def test_four_outputs_from_fake_model_run_enter_file_added_and_survive_reload(self):
        with tempfile.TemporaryDirectory() as root:
            service, _, stream = make_service(root)
            actor = access()
            cid = create(service, actor)
            calls = []

            def fake(key, payload, deadline, abort):
                calls.append(payload)
                if len(calls) == 1:
                    for i, (name, spec) in enumerate(SPECS.items()):
                        yield {
                            "type": "tool_delta",
                            "index": i,
                            "id": f"create-{i}",
                            "name": "write_file",
                            "arguments": json.dumps(
                                {"name": name, "content": spec}, ensure_ascii=False
                            ),
                        }
                    yield {"type": "finish", "reason": "tool_calls"}
                else:
                    yield {
                        "type": "text",
                        "text": "四个真实文件已生成，可在输出内容下载。",
                    }
                    yield {"type": "finish", "reason": "stop"}
                yield {
                    "type": "usage",
                    "prompt_tokens": 10,
                    "completion_tokens": 5,
                    "total_tokens": 15,
                }

            service.runner.stream_chat = fake
            try:
                request = body(service, actor, cid, mode="files")
                active = service.accept_run(actor, cid, request)
                events = "".join(service.stream(active))
                run = service.get_run(actor.principal, cid, active.run_id)
                self.assertEqual(run["status"], "completed", events[-2000:])
                self.assertEqual(run["tool_calls"], 4)
                self.assertEqual(run["model_requests"], 2)
                self.assertEqual(len(stream.calls), 0)
                scope = service.store.open_conversation(actor.principal, cid)
                entries = service.store.files(scope)
                self.assertEqual(len(entries), 4)
                for entry in entries:
                    validate(
                        self,
                        entry["name"],
                        service.store.read_file(scope, entry["id"])[1],
                    )
                self.assertEqual(events.count("event: file.added"), 4)
                before = len(calls)
                replay = service.accept_run(actor, cid, request)
                self.assertTrue(replay["replayed"])
                self.assertEqual(len(calls), before)
            finally:
                service.shutdown()
                service.store.close()
